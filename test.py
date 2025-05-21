import tkinter as tk
from tkinter import ttk, filedialog
from openpyxl import load_workbook

def select_excel_file():
    # Open file dialog to select an Excel file
    file_path = filedialog.askopenfilename(
        title="Select Excel File",
        filetypes=[("Excel files", "*.xlsx *.xls")]
    )
    print("File selected:", file_path)
    if file_path:
        try:
            workbook = load_workbook(file_path)
            sheet_names = workbook.sheetnames
            # Populate the dropdown with sheet names
            sheet_dropdown['values'] = sheet_names
            if sheet_names:
                sheet_dropdown.current(0)  # Select first sheet by default
        except Exception as e:
            print(f"Failed to load workbook: {e}")
            # Optionally: show a popup error dialog here

# Create the Tkinter app window
root = tk.Tk()
root.title("Excel File Loader")
root.geometry("300x150")

# Dropdown for sheet names
sheet_dropdown = ttk.Combobox(root, state="readonly")
sheet_dropdown.pack(pady=10)

# Button to select file
select_button = tk.Button(root, text="Select Excel File", command=select_excel_file)
select_button.pack(pady=20)

root.mainloop()
